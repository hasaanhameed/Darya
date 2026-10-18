import datetime as dt

from openpyxl import Workbook

from app.ingest.wapda.sheet_finder import day_sheets_in_workbook, find_day_sheets, parse_file_name

MONTHLY_SHEETS = ["Comparison", "Monthly", "Peak", "Summary"] + [str(d) for d in range(31, 0, -1)] + ["Sheet2", "Sheet5"]


def test_reads_month_year_from_name_variants():
    assert parse_file_name("JULY - 2022 (31).xlsx") == (2022, 7, 31)
    assert parse_file_name("FEB- 2026  (28).xls") == (2026, 2, 28)
    assert parse_file_name("April-2024 (30).xls") == (2024, 4, 30)
    assert parse_file_name("Daily data July- 2026 (25).xlsx") == (2026, 7, 25)
    assert parse_file_name("OCT- 2026 (07).xls") == (2026, 10, 7)


def test_skips_feb_29_to_31_in_non_leap_year():
    sheets = day_sheets_in_workbook("FEB- 2025 (28).xls", MONTHLY_SHEETS)
    assert len(sheets) == 28
    assert {"29", "30", "31"}.isdisjoint(s.sheet for s in sheets)


def test_keeps_feb_29_in_leap_year():
    sheets = day_sheets_in_workbook("Feb-2024 (29).xls", MONTHLY_SHEETS)
    assert dt.date(2024, 2, 29) in [s.expected_date for s in sheets]


def test_skips_sheet_31_in_30_day_month():
    sheets = day_sheets_in_workbook("SEP- 2024 (30).xls", MONTHLY_SHEETS)
    assert len(sheets) == 30
    assert "31" not in [s.sheet for s in sheets]


def test_ignores_summary_and_template_sheets():
    sheets = day_sheets_in_workbook("JULY - 2022 (31).xlsx", MONTHLY_SHEETS)
    assert all(s.sheet.isdigit() for s in sheets)


def test_expected_date_comes_from_file_name_and_sheet_name():
    sheets = day_sheets_in_workbook("AUG- 2024 (31).xls", MONTHLY_SHEETS)
    sheet_15 = next(s for s in sheets if s.sheet == "15")
    assert sheet_15.expected_date == dt.date(2024, 8, 15)


def test_single_day_workbook_reads_synopsis_with_day_from_name():
    sheet_names = ["Main", "Synopsis", "Ministry", "Hydel", "Chairman", "MP"]
    sheets = day_sheets_in_workbook("Daily data July- 2026 (27).xlsx", sheet_names)
    assert [(s.sheet, s.expected_date) for s in sheets] == [("Synopsis", dt.date(2026, 7, 27))]


def test_reports_unreadable_file_and_continues(tmp_path):
    workbook = Workbook()
    workbook.active.title = "Summary"
    workbook.create_sheet("1")
    workbook.create_sheet("2")
    workbook.save(tmp_path / "JAN- 2025 (31).xlsx")
    (tmp_path / "AUG- 2023 (31).xls").write_bytes(b"not a real workbook")

    sheets, unreadable = find_day_sheets(tmp_path)

    assert [(s.path.name, s.sheet) for s in sheets] == [("JAN- 2025 (31).xlsx", "1"), ("JAN- 2025 (31).xlsx", "2")]
    assert [p.name for p in unreadable] == ["AUG- 2023 (31).xls"]
